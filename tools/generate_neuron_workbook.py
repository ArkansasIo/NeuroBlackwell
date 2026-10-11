#!/usr/bin/env python3
"""Create an Excel workbook containing neuron, synapse, region, and provenance tables."""
import csv
from pathlib import Path
try:
    from openpyxl import Workbook
    from openpyxl.worksheet.table import Table, TableStyleInfo
except ImportError as exc:
    raise SystemExit("Install openpyxl to generate the workbook") from exc

ROOT=Path(__file__).resolve().parents[1]
src=ROOT/"brain3d/data/neuron_table.csv"
out=ROOT/"brain3d/data/NeuroBlackwell_Neuron_Data.xlsx"
wb=Workbook()
ws=wb.active; ws.title="Neurons"
with src.open(newline="",encoding="utf-8") as f:
    for row in csv.reader(f): ws.append(row)
for sheet,headers,rows in [
("Synapses",["source_neuron","target_neuron","weight","delay_ms","status"],[]),
("Regions",["region_id","name","parent","status"],[
("cortex.left","Left cerebral cortex","", "reference_scaffold"),
("cortex.right","Right cerebral cortex","", "reference_scaffold"),
("thalamus","Thalamus","", "reference_scaffold"),
("hippocampus.left","Left hippocampus","", "reference_scaffold"),
("hippocampus.right","Right hippocampus","", "reference_scaffold"),
("cerebellum","Cerebellum","", "reference_scaffold"),
("brainstem","Brainstem","", "reference_scaffold")]),
("Provenance",["dataset_id","modality","source","license","coordinate_system","status"],[
("NB-DEMO-BRAIN-001","synthetic","NeuroBlackwell","project-generated","normalized_brain_xyz","synthetic_demo")])]:
    s=wb.create_sheet(sheet); s.append(headers)
    for row in rows:s.append(row)
for s in wb.worksheets:
    if s.max_row>=2:
        ref=f"A1:{chr(64+s.max_column)}{s.max_row}"
        tab=Table(displayName=s.title.replace(" ","_")+"Table",ref=ref)
        tab.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showRowStripes=True)
        s.add_table(tab)
    s.freeze_panes="A2"
    for col in s.columns:
        s.column_dimensions[col[0].column_letter].width=min(36,max(12,max(len(str(c.value or "")) for c in col)+2))
wb.save(out)
print(out)
