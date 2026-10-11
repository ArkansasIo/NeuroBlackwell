from pathlib import Path
import json
from openpyxl import Workbook
from openpyxl.worksheet.table import Table,TableStyleInfo
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'data'/'NeuroBlackwell_Curriculum.xlsx'
def table(ws,name):
 if ws.max_row<2:return
 ref=f'A1:{chr(64+ws.max_column)}{ws.max_row}'; t=Table(displayName=name,ref=ref); t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True); ws.add_table(t); ws.freeze_panes='A2'
def main():
 c=json.loads((ROOT/'data/curriculum.json').read_text()); s=json.loads((ROOT/'data/subjects.json').read_text()); q=json.loads((ROOT/'data/question_bank.json').read_text())['questions']; w=Workbook(); ws=w.active; ws.title='Levels'; ws.append(['id','ordinal','difficulty_min','difficulty_max','focus'])
 for x in c['levels']: ws.append([x['id'],x['ordinal'],x['difficulty'][0],x['difficulty'][1],', '.join(x['focus'])])
 table(ws,'CurriculumLevels'); ws=w.create_sheet('Subjects'); ws.append(['id','name','domains'])
 for x in s['subjects']: ws.append([x['id'],x['name'],', '.join(x['domains'])])
 table(ws,'Subjects'); ws=w.create_sheet('Questions'); ws.append(list(q[0].keys()))
 for x in q: ws.append([', '.join(v) if isinstance(v,list) else v for v in x.values()])
 table(ws,'Questions'); ws=w.create_sheet('Readme'); ws['A1']='NeuroBlackwell Learning Benchmark'; ws['A2']='Synthetic assessment data; not a clinical IQ instrument.'; w.save(OUT); return OUT
if __name__=='__main__': main()
